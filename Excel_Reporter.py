# Excel_Reporter.py
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment

def export_audit_to_excel_with_chart(audit_df):
    """
    Creates an in-memory Excel workbook containing:
    1. Filtered transaction audit worklist.
    2. A status count breakdown table.
    3. An embedded native Pie/Circle Chart showing discrepancy distribution.
    """
    wb = Workbook()
    
    # ------------------ SHEET 1: AUDIT DATA ------------------
    ws_data = wb.active
    ws_data.title = "Audit_Worklist"

    # Convert Timestamp to string for clean export
    export_df = audit_df.copy()
    export_df["Timestamp"] = export_df["Timestamp"].dt.strftime("%Y-%m-%d %H:%M:%S")

    for r_idx, row in enumerate(dataframe_to_rows(export_df, index=False, header=True), 1):
        ws_data.append(row)
        if r_idx == 1:
            for cell in ws_data[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    # Auto-adjust column widths
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws_data.column_dimensions[col_letter].width = max(max_len + 2, 12)

    # ------------------ SHEET 2: SUMMARY & CHART ------------------
    ws_chart = wb.create_sheet(title="Executive_Summary")
    
    # Header block
    ws_chart["A1"] = "ReconLedger Audit Executive Summary"
    ws_chart["A1"].font = Font(size=14, bold=True, color="1F4E79")

    # Aggregate status counts
    counts = audit_df["Audit_Status"].value_counts().reset_index()
    counts.columns = ["Audit_Status", "Total_Count"]

    ws_chart["A3"] = "Audit Status"
    ws_chart["B3"] = "Transaction Count"
    ws_chart["A3"].font = Font(bold=True)
    ws_chart["B3"].font = Font(bold=True)

    for i, row in counts.iterrows():
        ws_chart[f"A{4 + i}"] = row["Audit_Status"]
        ws_chart[f"B{4 + i}"] = int(row["Total_Count"])

    row_end = 3 + len(counts)

    # Build the Native Excel Pie Chart
    pie = PieChart()
    pie.title = "Reconciliation Status Breakdown"
    labels = Reference(ws_chart, min_col=1, min_row=4, max_row=row_end)
    data = Reference(ws_chart, min_col=2, min_row=3, max_row=row_end)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.width = 16
    pie.height = 10

    # Place chart into sheet
    ws_chart.add_chart(pie, "D3")

    # Save to memory buffer
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer.getvalue()